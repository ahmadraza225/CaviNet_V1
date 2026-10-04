"""Read PatientIndex.xlsx (section 7.1).

Sheets `TB` and `NTM`; header row with Number, Gender (1 = male, 2 = female), Age, the 12
symptoms (1 = present, blank = absent), Annotation Status and (NTM only) Strain Type. Numbers
are written like the folders (`TB_001`, `Case_001`). Rows whose Number is not a case number,
such as the legend at the end of the TB sheet ("Gender 1:Male 2:Female"), are ignored and
listed. Annotation Status and Strain Type are not used (section 7.3).

The published file lists two NTM numbers twice with different values (and leaves two others
out). A case listed more than once is not guessed: its rows are reported and its clinical
values are left blank.
"""

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, BinaryIO

from cavinet_ml.dataset.source import DatasetError, canonical_case_id, class_of_case_id

SYMPTOMS = (
    "Chest pain",
    "Cough",
    "Expectoration",
    "Fever",
    "Chest tightness",
    "Haemoptysis",
    "Gasp",
    "Dyspnoea",
    "Chills",
    "Fatigue",
    "Night sweats",
    "Weight loss",
)


def column_name(label: str) -> str:
    """Manifest column for a spreadsheet label: "Night sweats" -> "night_sweats"."""
    return "_".join(label.lower().split())


SYMPTOM_COLUMNS = tuple(column_name(s) for s in SYMPTOMS)

_SYNONYMS = {
    "sex": "gender",
    "no": "number",
    "no.": "number",
    "hemoptysis": "haemoptysis",
    "dyspnea": "dyspnoea",
    "night sweat": "night sweats",
    "weight-loss": "weight loss",
}
_HEADER_SEARCH_ROWS = 10


@dataclass
class IndexRecord:
    case_id: str
    sheet: str  # TB or NTM
    row: int  # spreadsheet row number (1 = first row)
    sex: str | None  # "male" / "female"
    age: int | None
    symptoms: dict[str, int]  # SYMPTOM_COLUMNS -> 0 / 1
    problems: list[str] = field(default_factory=list)

    def values(self) -> tuple[Any, ...]:
        return (self.sex, self.age, tuple(sorted(self.symptoms.items())))


@dataclass
class PatientIndex:
    records: dict[str, IndexRecord]  # one unambiguous row per case
    duplicates: dict[str, list[IndexRecord]]  # listed more than once with different values
    ignored_rows: list[str]  # legend and note rows
    problems: list[str]


def _normalise(label: Any) -> str:
    text = " ".join(str(label).strip().lower().split()) if label is not None else ""
    return _SYNONYMS.get(text, text)


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _symptom(value: Any) -> int | None:
    if _blank(value):
        return 0
    text = str(value).strip()
    if text in ("1", "1.0") or value is True:
        return 1
    if text in ("0", "0.0") or value is False:
        return 0
    return None


def _sex(value: Any) -> str | None:
    text = "" if _blank(value) else str(value).strip().lower()
    if text in ("1", "1.0", "m", "male"):
        return "male"
    if text in ("2", "2.0", "f", "female"):
        return "female"
    return None


def _age(value: Any) -> int | None:
    if isinstance(value, bool) or _blank(value):
        return None
    try:
        number = float(str(value).strip())
    except ValueError:
        return None
    return int(number) if number.is_integer() and 0 <= number < 130 else None


def _header(rows: list[tuple[Any, ...]], sheet: str) -> tuple[int, dict[str, int]]:
    for i, row in enumerate(rows[:_HEADER_SEARCH_ROWS]):
        names = [_normalise(cell) for cell in row]
        if "number" in names:
            columns = {name: c for c, name in enumerate(names) if name}
            wanted = ["number", "gender", "age", *(_normalise(s) for s in SYMPTOMS)]
            missing = [w for w in wanted if w not in columns]
            if missing:
                raise DatasetError(
                    f"PatientIndex.xlsx sheet {sheet}: missing columns {', '.join(missing)}."
                )
            return i, columns
    raise DatasetError(f"PatientIndex.xlsx sheet {sheet}: no header row with a Number column.")


def _sheet_records(
    sheet: str, rows: list[tuple[Any, ...]], ignored: list[str], problems: list[str]
) -> list[IndexRecord]:
    header_at, columns = _header(rows, sheet)
    records = []
    for i, row in enumerate(rows[header_at + 1 :], start=header_at + 2):

        def cell(name: str, row: tuple[Any, ...] = row) -> Any:
            c = columns[_normalise(name)]
            return row[c] if c < len(row) else None

        if all(_blank(v) for v in row):
            continue
        number = cell("number")
        case_id = None if _blank(number) else canonical_case_id(str(number))
        if case_id is None:
            text = " | ".join(" ".join(str(v).split()) for v in row if not _blank(v))
            ignored.append(f"{sheet} row {i}: {text}")
            continue
        if class_of_case_id(case_id) != sheet:
            problems.append(f"{sheet} row {i}: {case_id} is listed in the {sheet} sheet; ignored")
            continue
        record = IndexRecord(case_id, sheet, i, _sex(cell("gender")), _age(cell("age")), {})
        if record.sex is None:
            record.problems.append(f"Gender {cell('gender')!r} is not 1 or 2")
        if record.age is None:
            record.problems.append(f"Age {cell('age')!r} is not a number")
        for label, column in zip(SYMPTOMS, SYMPTOM_COLUMNS, strict=True):
            value = _symptom(cell(label))
            if value is None:
                record.problems.append(f"{label} {cell(label)!r} is not 1 or blank; read as absent")
                value = 0
            record.symptoms[column] = value
        problems.extend(f"{sheet} row {i} ({case_id}): {p}" for p in record.problems)
        records.append(record)
    return records


def read_patient_index(source: Path | str | BinaryIO) -> PatientIndex:
    from openpyxl import load_workbook  # training extra; imported only when needed

    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001 - zip, XML and format errors alike
        raise DatasetError(f"PatientIndex.xlsx cannot be read ({error}).") from error
    sheets = {ws.title.strip().upper(): ws for ws in workbook.worksheets}
    missing = [name for name in ("TB", "NTM") if name not in sheets]
    if missing:
        raise DatasetError(f"PatientIndex.xlsx has no {' or '.join(missing)} sheet.")

    ignored: list[str] = []
    problems: list[str] = []
    by_case: dict[str, list[IndexRecord]] = {}
    for name in ("TB", "NTM"):
        rows = [tuple(r) for r in sheets[name].iter_rows(values_only=True)]
        for record in _sheet_records(name, rows, ignored, problems):
            by_case.setdefault(record.case_id, []).append(record)
    workbook.close()

    records: dict[str, IndexRecord] = {}
    duplicates: dict[str, list[IndexRecord]] = {}
    for case_id, found in by_case.items():
        if len({r.values() for r in found}) == 1:
            records[case_id] = found[0]
            if len(found) > 1:
                rows = ", ".join(str(r.row) for r in found)
                problems.append(
                    f"{case_id} is listed {len(found)} times (rows {rows}), identically"
                )
        else:
            duplicates[case_id] = found
            rows = ", ".join(str(r.row) for r in found)
            problems.append(
                f"{case_id} is listed {len(found)} times with different values (rows {rows}); "
                "its clinical values are left blank"
            )
    return PatientIndex(records, duplicates, ignored, problems)
